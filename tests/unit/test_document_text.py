"""Extraction must cover every format the indexer accepts.

Uploads accept PDF, HWP, HWPX, DOCX and PPTX, and all five are in
`rag_index_plan.INDEXABLE_EXTS`, but extraction only handled `.pdf` and
`.docx`; everything else fell through to `read_text(utf-8, errors="ignore")`.

`.pptx` and `.hwpx` are zip files and `.hwp` is a binary compound document.
Read as UTF-8 they raise nothing and quietly put broken strings into the index,
so search does not fail — it returns garbage. Those chunks then become KBD
probes and produce a coverage number.

The extractor imports its parsers lazily, so the parts the standard library
can check (HWPX, plain text, unsupported formats, dispatch) always run and the
rest uses `importorskip`.
"""

from __future__ import annotations

import ast
import pathlib
import zipfile

import pytest

from server.core.document_text import (
    BINARY_EXTS,
    PAGED_EXTS,
    UNSUPPORTED_EXTS,
    extract_pages,
    hwpx_text,
    is_unsupported,
    unsupported_reason,
)
from server.core.rag_index_plan import INDEXABLE_EXTS

REPO = pathlib.Path(__file__).resolve().parents[2]
DATA_TASKS = REPO / "celery_app/tasks/data_tasks.py"

_HWPX_NS = (
    'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" xmlns:tbl="http://www.hancom.co.kr/hwpml/2011/table"'
)


def _make_hwpx(path: pathlib.Path, body: str) -> pathlib.Path:
    """HWPX is zip + OWPML XML, so it can be synthesised with the standard library."""
    xml = f'<?xml version="1.0" encoding="UTF-8"?>\n<hp:sec {_HWPX_NS}>{body}</hp:sec>'
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("Contents/section0.xml", xml)
    return path


# ══════════════════════════════════════════════════════════════════
# 1. HWPX — zip + XML, fully checkable with the standard library
# ══════════════════════════════════════════════════════════════════


class TestHwpx:
    def test_extracts_paragraphs(self, tmp_path):
        f = _make_hwpx(
            tmp_path / "a.hwpx",
            "<hp:para>KISTI operates national science and technology data.</hp:para>",
        )
        assert "national science and technology data" in hwpx_text(f)

    def test_extracts_table_cells(self, tmp_path):
        f = _make_hwpx(
            tmp_path / "t.hwpx",
            "<tbl:tbl><tbl:tr>"
            "<tbl:tc><hp:para>coverage</hp:para></tbl:tc>"
            "<tbl:tc><hp:para>23.3%</hp:para></tbl:tc>"
            "</tbl:tr></tbl:tbl>",
        )
        text = hwpx_text(f)
        assert "coverage" in text and "23.3%" in text

    def test_short_paragraph_becomes_a_heading(self, tmp_path):
        """Preserves the original parser rule: short and unpunctuated is a heading."""
        f = _make_hwpx(tmp_path / "h.hwpx", "<hp:para>Study overview</hp:para>")
        assert hwpx_text(f).startswith("## Study overview")

    def test_binary_is_not_read_as_text(self, tmp_path):
        """What this fixes: the zip header 'PK\\x03\\x04' used to be indexed."""
        f = _make_hwpx(tmp_path / "b.hwpx", "<hp:para>Body text.</hp:para>")
        ((_, text),) = extract_pages(f)
        assert "PK" not in text[:4]
        assert "Body text." in text

    def test_page_numbers_are_not_invented(self, tmp_path):
        """HWPX has no page unit, and inventing one would make provenance a lie."""
        f = _make_hwpx(tmp_path / "p.hwpx", "<hp:para>Body</hp:para>")
        assert extract_pages(f)[0][0] is None

    def test_one_broken_xml_part_does_not_discard_the_document(self, tmp_path):
        f = tmp_path / "mixed.hwpx"
        with zipfile.ZipFile(f, "w") as zf:
            zf.writestr("Contents/broken.xml", "<<< not xml")
            zf.writestr(
                "Contents/section0.xml",
                f'<?xml version="1.0"?><hp:sec {_HWPX_NS}><hp:para>Surviving body.</hp:para></hp:sec>',
            )
        assert "Surviving body." in hwpx_text(f)

    def test_a_non_zip_gives_an_empty_result(self, tmp_path):
        """It must not raise and kill the whole indexing run."""
        f = tmp_path / "fake.hwpx"
        f.write_bytes(b"not a zip")
        assert extract_pages(f) == []


# ══════════════════════════════════════════════════════════════════
# 2. Unopenable formats — never passed over silently
# ══════════════════════════════════════════════════════════════════


class TestUnsupported:
    def test_hwp_is_unsupported(self):
        """A binary compound document needs an office converter, which the worker lacks."""
        assert is_unsupported(".hwp") is True

    def test_returns_an_empty_result(self, tmp_path):
        """Indexing nothing beats indexing a broken string."""
        f = tmp_path / "x.hwp"
        f.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)
        assert extract_pages(f) == []

    def test_the_reason_is_never_empty(self):
        assert unsupported_reason(".hwp", "en").strip()

    def test_the_reason_is_english(self):
        assert not any("\uac00" <= c <= "\ud7a3" for c in unsupported_reason(".hwp", "en"))

    def test_supported_formats_have_no_reason(self):
        assert unsupported_reason(".pdf") == ""


# ══════════════════════════════════════════════════════════════════
# 3. Plain text and dispatch
# ══════════════════════════════════════════════════════════════════


class TestTextLike:
    @pytest.mark.parametrize("name", ["a.txt", "a.md", "a.json", "a.jsonl"])
    def test_plain_text_is_read_as_is(self, tmp_path, name):
        f = tmp_path / name
        f.write_text("body content", encoding="utf-8")
        assert extract_pages(f) == [(None, "body content")]

    def test_broken_encoding_does_not_raise(self, tmp_path):
        f = tmp_path / "a.txt"
        f.write_bytes(b"\xff\xfe ok")
        assert extract_pages(f)[0][1].endswith("ok")


class TestContract:
    def test_every_indexable_format_is_handled(self):
        """A format in `INDEXABLE_EXTS` with no rule here falls back to plain text,
        which is the defect this module exists to fix."""
        handled = BINARY_EXTS | set(UNSUPPORTED_EXTS) | {".txt", ".md", ".json", ".jsonl"}
        assert INDEXABLE_EXTS <= handled, sorted(INDEXABLE_EXTS - handled)

    def test_only_paged_formats_carry_numbers(self):
        """PDFs have pages and PPTX has slides. DOCX and HWPX have neither."""
        assert PAGED_EXTS == {".pdf", ".pptx"}
        assert ".docx" not in PAGED_EXTS and ".hwpx" not in PAGED_EXTS

    def test_unsupported_and_handled_do_not_overlap(self):
        assert not (BINARY_EXTS & set(UNSUPPORTED_EXTS))


# ══════════════════════════════════════════════════════════════════
# 4. Only when the libraries are installed (not a CI dependency)
# ══════════════════════════════════════════════════════════════════


class TestPptx:
    def test_one_entry_per_slide_numbered_from_one(self, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation

        prs = Presentation()
        for title in ("First slide", "Second slide"):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = title
        f = tmp_path / "a.pptx"
        prs.save(str(f))

        pages = extract_pages(f)
        assert [p for p, _ in pages] == [1, 2]
        assert "First slide" in pages[0][1]
        assert "Second slide" in pages[1][1]

    def test_extracts_table_cells(self, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])
        t = s.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        t.cell(0, 0).text = "coverage"
        t.cell(0, 1).text = "23.3%"
        f = tmp_path / "t.pptx"
        prs.save(str(f))

        text = extract_pages(f)[0][1]
        assert "coverage" in text and "23.3%" in text

    def test_binary_is_not_read_as_text(self, tmp_path):
        pytest.importorskip("pptx")
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Body"
        f = tmp_path / "b.pptx"
        prs.save(str(f))
        assert "PK" not in extract_pages(f)[0][1][:4]


class TestWiring:
    @pytest.fixture(scope="class")
    def src(self):
        return DATA_TASKS.read_text(encoding="utf-8")

    def test_the_indexing_task_uses_this_module(self, src):
        assert "from server.core.document_text import extract_pages" in src

    def test_extraction_is_not_reimplemented_in_the_task(self, src):
        """A copy would drift from this one."""
        tree = ast.parse(src)
        local = [
            n.name
            for n in ast.walk(tree)
            if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in {"_extract_pages", "_extract_text"}
        ]
        assert not local, f"extraction logic reappeared inside the task: {local}"

    def test_the_plain_text_fallback_does_not_return(self, src):
        """Accepting every format through `read_text(..., errors="ignore")` was the defect."""
        tree = ast.parse(src)
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and ast.unparse(node.func).endswith("read_text"):
                rendered = ast.unparse(node)
                assert "errors='ignore'" not in rendered, rendered
